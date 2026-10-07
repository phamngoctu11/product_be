from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "THIET_KE_CODE_VA_ENDPOINT.xlsx"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")

STATUS_COLORS = {
    "GIỮ": GREEN,
    "SỬA": YELLOW,
    "TẠO MỚI": LIGHT_BLUE,
    "ĐÃ CÓ KHUNG": ORANGE,
    "MIGRATE": ORANGE,
    "LOẠI SAU MIGRATION": RED,
}


def title(ws, text: str, subtitle: str | None = None, columns: int = 8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    c = ws.cell(1, 1, text)
    c.font = Font(name="Arial", size=18, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
        c = ws.cell(2, 1, subtitle)
        c.font = Font(name="Arial", size=10, italic=True, color="404040")
        c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 34


def paragraph(ws, row: int, text: str, columns: int = 8, fill: str = WHITE):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    c = ws.cell(row, 1, text)
    c.font = Font(name="Arial", size=10)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = max(34, min(100, 18 * (text.count("\n") + 2)))


def section(ws, row: int, text: str, columns: int = 8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    c = ws.cell(row, 1, text)
    c.font = Font(name="Arial", size=12, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24


def table(ws, start_row: int, headers: list[str], rows: list[list], name: str, widths=None):
    for col, value in enumerate(headers, 1):
        c = ws.cell(start_row, col, value)
        c.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for r_idx, values in enumerate(rows, start_row + 1):
        for c_idx, value in enumerate(values, 1):
            c = ws.cell(r_idx, c_idx, value)
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if r_idx % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F8FBFD")
            if isinstance(value, str) and value in STATUS_COLORS:
                c.fill = PatternFill("solid", fgColor=STATUS_COLORS[value])
                c.font = Font(name="Arial", size=9, bold=True)
        ws.row_dimensions[r_idx].height = 42
    end_row = start_row + len(rows)
    if rows:
        ref = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
        tab = Table(displayName=name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False)
        ws.add_table(tab)
    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{max(start_row, end_row)}"
    if widths:
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    else:
        for i, header in enumerate(headers, 1):
            sample = [str(header)] + [str(r[i - 1]) if i - 1 < len(r) else "" for r in rows[:100]]
            ws.column_dimensions[get_column_letter(i)].width = min(55, max(10, max(len(x) for x in sample) + 2))
    return end_row


def add_sheet(wb, name: str, heading: str, subtitle: str, headers: list[str], rows: list[list], widths=None):
    ws = wb.create_sheet(name)
    title(ws, heading, subtitle, len(headers))
    table(ws, 4, headers, rows, "T_" + name.replace("-", "_").replace(" ", "_"), widths)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


objects = []


def obj(kind, filename, package, status, phase, workflows, purpose, dependencies=""):
    objects.append([kind, filename, package, status, phase, workflows, purpose, dependencies])


# Entity and persistence files. One public Java type per file; file name equals the contained type.
for data in [
    ("Entity", "Order.java", "entity", "SỬA", "02,04-12", "WF01-WF09", "Aggregate gốc của vòng đời đơn; trạng thái, payment, staff, contact snapshot và các mốc thời gian.", "OrderItem, OrderAssignment"),
    ("Entity", "OrderItem.java", "entity", "SỬA", "02,04,05,10,12", "WF01,WF05,WF07,WF08", "Dòng hàng; lưu snapshot catalog/custom, đơn giá nullable, quantity và thời lượng sản xuất.", "Order, ProductVariant"),
    ("Entity", "OrderContactSnapshot.java", "entity", "GIỮ", "04", "WF02,WF08,WF09", "Snapshot họ tên, email, điện thoại, địa chỉ và ghi chú tại thời điểm đặt.", "Order"),
    ("Entity", "OrderStatusHistory.java", "entity", "SỬA", "03,06-12", "WF02-WF09", "Lịch sử trạng thái hiển thị; actor, lý do, mốc thời gian và tham chiếu nghiệp vụ.", "OrderTransitionService"),
    ("Entity", "CustomRequest.java", "entity", "ĐÃ CÓ KHUNG", "05", "WF01", "Nhiều bản nháp custom của USER; snapshot nội dung/ảnh khi submit thành Order.", "User, Order"),
    ("Entity", "OrderAssignment.java", "entity", "ĐÃ CÓ KHUNG", "07", "WF03,WF05,WF09", "Một staff phụ trách một Order; unique assignment hoạt động theo staff và Order.", "Order, User(STAFF)"),
    ("Entity", "ChatThread.java", "entity", "TẠO MỚI", "08", "WF03", "Phòng chat riêng theo Order, chỉ USER sở hữu và STAFF phụ trách là thành viên.", "Order, OrderAssignment"),
    ("Entity", "ChatMessage.java", "entity", "MIGRATE", "08", "WF03", "Tin nhắn gắn ChatThread/Order; lưu sender, nội dung, attachment và thời gian.", "ChatThread"),
    ("Entity", "OrderAgreement.java", "entity", "ĐÃ CÓ KHUNG", "08", "WF03", "Phiên bản form thỏa thuận; giá/deadline/spec và quyết định của staff.", "Order, UserVoucher"),
    ("Entity", "PaymentAttempt.java", "entity", "ĐÃ CÓ KHUNG", "09", "WF04", "Một lần tạo giao dịch online; hạn 1 giờ, kết quả provider và cờ đối soát.", "Order"),
    ("Entity", "ProductionCheckpoint.java", "entity", "SỬA", "10", "WF05", "Checkpoint theo loại OrderItem và attempt; INITIAL_SHAPE hoặc FINAL_QC.", "OrderItem"),
    ("Entity", "ProductionCheckpointImage.java", "entity", "SỬA", "10", "WF05", "Ảnh tiến trình thuộc đúng checkpoint/attempt.", "ProductionCheckpoint"),
    ("Entity", "ProductionDecision.java", "entity", "SỬA", "10", "WF05", "Quyết định KCS checkpoint 2: đạt hoặc yêu cầu làm lại, có lý do.", "ProductionCheckpoint"),
    ("Entity", "ChangeRequest.java", "entity", "ĐÃ CÓ KHUNG", "11", "WF06", "Yêu cầu đổi chi tiết chưa triển khai; không thay đổi giá/payment/deadline.", "Order, OrderItem"),
    ("Entity", "OrderReceipt.java", "entity", "ĐÃ CÓ KHUNG", "12", "WF08", "Snapshot số lượng thực nhận và xác nhận chênh lệch trước DELIVERED.", "Order"),
    ("Entity", "Product.java", "entity", "SỬA", "04", "Catalog", "Mẫu catalog, availability và madeDay do manager nhập, madeDay >= 2.", "ProductVariant"),
    ("Entity", "ProductVariant.java", "entity", "SỬA", "04", "Catalog", "Biến thể và giá catalog chính thức dùng để tạo snapshot OrderItem.", "Product"),
    ("Entity", "Cart.java", "entity", "SỬA", "04", "Checkout", "Giỏ theo USER hoặc guestSessionId.", "CartItem"),
    ("Entity", "CartItem.java", "entity", "SỬA", "04", "Checkout", "Variant và quantity dương; checkout có thể chọn một phần giỏ.", "Cart, ProductVariant"),
    ("Entity", "VoucherTemplate.java", "entity", "SỬA", "04,08", "Checkout,WF03", "Chính sách voucher, quota, hạn và minimum order.", "UserVoucher, GuestVoucherUsage"),
    ("Entity", "UserVoucher.java", "entity", "SỬA", "04,08", "Checkout,WF03", "Voucher ví USER, consume/restore nguyên tử.", "VoucherTemplate, User"),
    ("Entity", "GuestVoucherUsage.java", "entity", "SỬA", "04,06", "Checkout,WF09", "Quota guest và dấu vết chống spam theo session/email/phone hash.", "VoucherTemplate"),
    ("Entity", "Notification.java", "entity", "GIỮ", "03,04-12", "WF01-WF09", "Thông báo bền vững cho USER/STAFF; guest không có notification nội bộ.", "User"),
    ("Entity", "NotificationRead.java", "entity", "GIỮ", "03,04-12", "WF01-WF09", "Trạng thái đọc thông báo theo người nhận.", "Notification"),
    ("Entity", "OrderItemAssignment.java", "entity", "MIGRATE", "13", "Legacy", "Migrate khỏi phân công cấp item sang OrderAssignment cấp Order.", "OrderAssignment"),
    ("Entity", "InventoryTransaction.java", "entity", "LOẠI SAU MIGRATION", "13", "Ngoài phạm vi", "Không thuộc hệ thống made-to-order mới; chỉ xóa sau khi không còn consumer.", "")
]: obj(*data)

for name, status, phase, purpose in [
    ("OrderRepository.java", "SỬA", "03-12", "Khóa Order, query theo actor/status và optimistic version."),
    ("OrderItemRepository.java", "SỬA", "04,10,12", "Truy vấn toàn bộ dòng hàng và checkpoint/receipt theo Order."),
    ("CustomRequestRepository.java", "ĐÃ CÓ KHUNG", "05", "Draft theo owner, khóa submit và liên kết Order duy nhất."),
    ("OrderAssignmentRepository.java", "ĐÃ CÓ KHUNG", "07", "Khóa assignment, kiểm tra staff/order đang active và release."),
    ("ChatThreadRepository.java", "TẠO MỚI", "08", "Tìm hoặc tạo duy nhất thread theo orderId."),
    ("ChatMessageRepository.java", "SỬA", "08", "Phân trang message theo thread; không truy cập chỉ bằng userId."),
    ("OrderAgreementRepository.java", "ĐÃ CÓ KHUNG", "08", "Phiên bản agreement duy nhất theo order/version."),
    ("PaymentAttemptRepository.java", "ĐÃ CÓ KHUNG", "09", "Khóa attempt theo request/provider transaction và webhook idempotent."),
    ("ProductionCheckpointRepository.java", "SỬA", "10", "Checkpoint theo item/type/attempt, khóa review."),
    ("ProductionCheckpointImageRepository.java", "GIỮ", "10", "Ảnh theo checkpoint."),
    ("ProductionDecisionRepository.java", "SỬA", "10", "Quyết định manager theo checkpoint/version."),
    ("ChangeRequestRepository.java", "ĐÃ CÓ KHUNG", "11", "Yêu cầu đổi theo order và status."),
    ("OrderReceiptRepository.java", "ĐÃ CÓ KHUNG", "12", "Một receipt cuối cho Order."),
    ("CartRepository.java", "SỬA", "04", "Khóa giỏ USER/Guest trong checkout."),
    ("CartItemRepository.java", "SỬA", "04", "Xóa đúng các item đã checkout."),
    ("ProductRepository.java", "SỬA", "04", "Catalog availability/madeDay."),
    ("ProductVariantRepository.java", "SỬA", "04", "Khóa/đọc giá server và trạng thái variant."),
    ("UserVoucherRepository.java", "SỬA", "04,08", "Consume/restore voucher USER nguyên tử."),
    ("GuestVoucherUsageRepository.java", "SỬA", "04,06", "Quota guest nguyên tử, giữ anti-spam khi hoàn quota."),
]: obj("Repository", name, "repository", status, phase, "Theo phạm vi", purpose)

for name, status, phase, wf, purpose, dep in [
    ("ProductService.java", "SỬA", "04", "Catalog", "Quản lý catalog, madeDay và availability.", "ProductRepository"),
    ("CatalogDurationCalculator.java", "TẠO MỚI", "04", "Checkout", "Tính thời lượng từng OrderItem đúng công thức và làm tròn đã chốt.", "Không truy cập DB"),
    ("ProductionCalendarService.java", "TẠO MỚI", "04,09", "WF04", "Cộng số ngày sản xuất, không tính ngày bắt đầu, Chủ nhật và ngày nghỉ cấu hình.", "Holiday source"),
    ("CartService.java", "SỬA", "04", "Checkout", "Chỉ quản lý thêm/sửa/xóa/đọc giỏ; tách checkout ra service riêng.", "CartRepository"),
    ("CheckoutService.java", "TẠO MỚI", "04", "WF02 input", "Checkout USER/Guest, snapshot, voucher, PENDING_APPROVAL, idempotency/outbox.", "Cart, Voucher, Duration"),
    ("CustomRequestService.java", "TẠO MỚI", "05", "WF01", "CRUD draft và submit idempotent thành CUSTOM Order.", "CustomRequestRepository"),
    ("CancelOrderService.java", "TẠO MỚI", "06", "WF09", "Kernel hủy dùng chung, policy state, voucher, assignment và event.", "OrderTransitionService"),
    ("ManagerOrderService.java", "SỬA", "07", "WF02", "Manager duyệt khả năng làm hoặc từ chối; không quản trị chat.", "Assignment, Cancel"),
    ("OrderAssignmentService.java", "TẠO MỚI", "07", "WF03", "Assign/claim/release nguyên tử theo cùng một rule.", "OrderAssignmentRepository"),
    ("OrderAccessPolicy.java", "TẠO MỚI", "07-12", "WF02-WF09", "Quyền chủ đơn, staff phụ trách, manager; quyền order tách quyền chat.", "AuthService"),
    ("OrderChatService.java", "TẠO MỚI", "08", "WF03", "Tạo thread idempotent, membership USER–STAFF, message/history.", "ChatThreadRepository"),
    ("OrderAgreementService.java", "TẠO MỚI", "08", "WF03", "USER submit form phiên bản; STAFF xác nhận/yêu cầu sửa; dừng timer khi form hợp lệ.", "OrderAgreementRepository"),
    ("OrderPaymentService.java", "TẠO MỚI", "09", "WF04", "Mở payment sau ORDER_ACCEPTED, tạo attempt và hạn PT1H.", "PaymentAttemptRepository"),
    ("PaymentWebhookService.java", "TẠO MỚI", "09", "WF04", "Xác thực webhook, số tiền, idempotency; không tự ORDER_CREATING.", "MomoService"),
    ("MomoService.java", "SỬA", "09", "WF04", "Adapter provider, không sở hữu state machine.", "HTTP provider"),
    ("ProductionService.java", "TẠO MỚI", "09", "WF04", "STAFF bấm bắt đầu, chuyển ORDER_CREATING và tính deadline.", "Calendar, Transition"),
    ("ProductionCheckpointService.java", "TẠO MỚI", "10", "WF05", "Submit checkpoint, tiến trình, review/rework, hội tụ READY_TO_SHIP.", "Checkpoint repositories"),
    ("ChangeRequestService.java", "TẠO MỚI", "11", "WF06", "USER tạo, STAFF quyết định; chi tiết chưa làm, không đổi giá/deadline.", "ChangeRequestRepository"),
    ("OrderShippingService.java", "TẠO MỚI", "12", "WF07", "Ghi exportedQuantity và READY_TO_SHIP -> SHIPPING.", "OrderItemRepository"),
    ("OrderReceiptService.java", "TẠO MỚI", "12", "WF08", "Ghi receivedQuantity, xác nhận chênh lệch và DELIVERED.", "OrderReceiptRepository"),
    ("ReorderService.java", "TẠO MỚI", "12", "WF08", "Catalog thêm giỏ; custom tạo draft mới.", "Cart, CustomRequest"),
    ("OrderQueryService.java", "TẠO MỚI", "07-12", "WF02-WF09", "Read model danh sách/chi tiết/lịch sử theo quyền.", "Order repositories"),
    ("OrderService.java", "MIGRATE", "04-13", "Legacy", "Thu nhỏ dần; bỏ logic checkout/payment/inventory/checkpoint đã tách.", "Các service đích"),
]: obj("Service", name, "service", status, phase, wf, purpose, dep)

for name, status, phase, purpose in [
    ("ProductController.java", "SỬA", "04", "API catalog và madeDay/availability."),
    ("CartController.java", "SỬA", "04", "API giỏ; bỏ userId do client tự nhận và route approve legacy."),
    ("CheckoutController.java", "TẠO MỚI", "04", "Checkout USER từ JWT."),
    ("GuestCheckoutController.java", "SỬA", "04", "Checkout Guest COD."),
    ("CustomRequestController.java", "TẠO MỚI", "05", "Draft/submit custom của USER."),
    ("ManagerOrderController.java", "SỬA", "07,10", "Review, assignment và KCS; không có quyền chat."),
    ("StaffOrderController.java", "SỬA", "07,09,12", "Claim, danh sách của staff và bàn giao."),
    ("OrderChatController.java", "TẠO MỚI", "08", "Thông tin thread và lịch sử; realtime qua WebSocket."),
    ("OrderAgreementController.java", "TẠO MỚI", "08", "Submit/review agreement."),
    ("OrderPaymentController.java", "TẠO MỚI", "09", "Mở và đọc payment attempt."),
    ("MomoController.java", "SỬA", "09", "Webhook provider; route cũ tương thích có thời hạn."),
    ("ProductionController.java", "TẠO MỚI", "09,10", "Bắt đầu và checkpoint."),
    ("ChangeRequestController.java", "TẠO MỚI", "11", "ChangeRequest theo Order."),
    ("OrderController.java", "SỬA", "06,12", "Query, cancel, receipt, reorder của USER."),
    ("GuestOrderController.java", "TẠO MỚI", "06,12", "Lookup/cancel/receipt bằng token scope."),
]: obj("Controller", name, "controller", status, phase, "Theo phạm vi", purpose)

for name, status, phase, purpose in [
    ("OrderMapper.java", "SỬA", "04-12", "Map giá custom null, snapshot và lifecycle mới."),
    ("CheckoutMapper.java", "TẠO MỚI", "04", "Map kết quả checkout; không tính tiền trong mapper."),
    ("CustomRequestMapper.java", "TẠO MỚI", "05", "Map draft/detail."),
    ("OrderAgreementMapper.java", "TẠO MỚI", "08", "Map agreement/version."),
    ("PaymentAttemptMapper.java", "TẠO MỚI", "09", "Ẩn provider payload nhạy cảm."),
    ("ProductionCheckpointMapper.java", "TẠO MỚI", "10", "Map checkpoint, ảnh, decision."),
    ("OrderReceiptMapper.java", "TẠO MỚI", "12", "Map receipt/mismatch."),
]: obj("Mapper", name, "mapper", status, phase, "Theo phạm vi", purpose)


dto_rows = [
    ["CheckoutRequest.java", "Request", "TẠO MỚI", "04", "selectedVariantIds, userVoucherId, paymentMethod, note", "USER checkout; actor lấy từ JWT."],
    ["GuestCheckoutRequest.java", "Request", "SỬA", "04", "selectedVariantIds, contact, guestVoucherCode, note", "Guest chỉ CATALOG + COD."],
    ["CheckoutResponseDTO.java", "Response", "SỬA", "04", "orderId, status, version, totals, paymentMethod", "Không trả payment URL trước review."],
    ["CreateCustomRequest.java", "Request", "TẠO MỚI", "05", "name, description, quantity, images/spec", "Tạo draft."],
    ["UpdateCustomRequest.java", "Request", "TẠO MỚI", "05", "version và trường draft", "Chỉ draft owner."],
    ["SubmitCustomRequest.java", "Request", "TẠO MỚI", "05", "expectedVersion, contact/note", "Idempotency-Key ở header."],
    ["CustomRequestDTO.java", "Response", "TẠO MỚI", "05", "draft, status, orderId, version", "Không lộ draft user khác."],
    ["ReviewOrderRequest.java", "Request", "TẠO MỚI", "07", "decision, reason, optional staffId, expectedVersion", "Manager review."],
    ["AssignOrderRequest.java", "Request", "TẠO MỚI", "07", "staffId, expectedVersion", "Manager assign."],
    ["OrderAssignmentDTO.java", "Response", "TẠO MỚI", "07", "orderId, staffId, source, status, timestamps", "Assignment cấp Order."],
    ["ChatThreadDTO.java", "Response", "TẠO MỚI", "08", "threadId, orderId, peer, createdAt", "Không chứa manager member."],
    ["SendOrderMessageRequest.java", "Request", "TẠO MỚI", "08", "clientMessageId, content, attachments", "Idempotent theo clientMessageId."],
    ["OrderChatMessageDTO.java", "Response", "TẠO MỚI", "08", "messageId, sender, content, sentAt", "Cho member hợp lệ."],
    ["SubmitAgreementRequest.java", "Request", "TẠO MỚI", "08", "spec, amount/deadline, voucherId, expectedVersion", "Catalog price read-only; custom amount USER nhập theo trao đổi."],
    ["ReviewAgreementRequest.java", "Request", "TẠO MỚI", "08", "agreementVersion, decision, reason, expectedOrderVersion", "Staff phụ trách."],
    ["OrderAgreementDTO.java", "Response", "TẠO MỚI", "08", "version, snapshot, totals, status, timestamps", "Giá custom null trước chốt."],
    ["CreatePaymentAttemptRequest.java", "Request", "TẠO MỚI", "09", "expectedOrderVersion", "ONLINE + ORDER_ACCEPTED."],
    ["PaymentAttemptDTO.java", "Response", "TẠO MỚI", "09", "attemptId, status, amount, dueAt, payUrl", "Không trả provider secret/result raw."],
    ["MomoWebhookRequest.java", "Request", "TẠO MỚI", "09", "provider fields/signature", "Validate chữ ký và số tiền."],
    ["StartProductionRequest.java", "Request", "TẠO MỚI", "09", "expectedVersion, plannedStartReason?", "Staff bấm bắt đầu."],
    ["SubmitCheckpointRequest.java", "Request", "TẠO MỚI", "10", "orderItemId, type, attempt, images, note", "Checkpoint theo loại, không theo từng chiếc."],
    ["ReviewCheckpointRequest.java", "Request", "TẠO MỚI", "10", "decision, reason, expectedCheckpointVersion", "Chỉ checkpoint 2 qua manager."],
    ["CreateChangeRequest.java", "Request", "TẠO MỚI", "11", "orderItemId, requestedChanges, expectedVersion", "Không có price/deadline change."],
    ["ReviewChangeRequest.java", "Request", "TẠO MỚI", "11", "decision, reason, expectedVersion", "Staff phụ trách."],
    ["ChangeRequestDTO.java", "Response", "TẠO MỚI", "11", "request, decision, timestamps", "Không phát sinh payment."],
    ["ShipOrderRequest.java", "Request", "TẠO MỚI", "12", "items[orderItemId, exportedQuantity], note, expectedVersion", "Một snapshot bàn giao."],
    ["ReceiptConfirmRequest.java", "Request", "SỬA", "12", "items[orderItemId, receivedQuantity], acceptMismatch, expectedVersion", "USER/Guest dùng cùng nghiệp vụ."],
    ["OrderReceiptDTO.java", "Response", "TẠO MỚI", "12", "items snapshot, mismatch, confirmedAt", "Kết quả nhận hàng."],
    ["CancelOrderRequest.java", "Request", "TẠO MỚI", "06", "reason, expectedVersion", "Chỉ trước ORDER_ACCEPTED cho USER/Guest."],
    ["ReorderRequest.java", "Request", "TẠO MỚI", "12", "expectedSourceOrderVersion", "Catalog thêm giỏ; custom tạo draft."],
    ["OrderActionResponse.java", "Response", "TẠO MỚI", "06-12", "orderId, status, paymentStatus, version, message", "Response chung cho command."],
    ["OrderDTO.java", "Response", "SỬA", "04-12", "Lifecycle, payment, snapshots, assigned staff, deadlines", "Giá custom chưa chốt là null."],
    ["OrderItemDTO.java", "Response", "SỬA", "04-12", "sourceType, names/spec/price/duration/completion", "Dữ liệu snapshot."],
]


endpoint_rows = [
    ["CAT-01", "Catalog", "GET", "/api/products", "PUBLIC", "Tra cứu danh sách catalog/filter; chỉ hiển thị dữ liệu công khai.", "Query params", "ProductDTO page", "ProductService", "-", "GIỮ"],
    ["CAT-02", "Catalog", "GET", "/api/products/{id}", "PUBLIC", "Xem chi tiết catalog và madeDay tham khảo.", "path id", "ProductDTO", "ProductService", "-", "GIỮ"],
    ["CAT-03", "Catalog", "POST", "/api/products", "MANAGER", "Tạo catalog; bắt buộc madeDay hữu hạn >= 2.", "CreateProductRequest", "ProductDTO", "ProductService", "Idempotency-Key khuyến nghị", "SỬA"],
    ["CAT-04", "Catalog", "PUT", "/api/products/{id}", "MANAGER", "Cập nhật catalog; giá mới không sửa ngược OrderItem snapshot.", "UpdateProductRequest", "ProductDTO", "ProductService", "expectedVersion", "SỬA"],
    ["CART-01", "Cart", "GET", "/api/cart", "USER/GUEST", "Đọc giỏ theo JWT hoặc guestSessionId.", "headers", "CartResDTO", "CartService", "-", "GIỮ"],
    ["CART-02", "Cart", "POST", "/api/cart/items", "USER/GUEST", "Thêm variant ACCEPTING_ORDERS và quantity dương.", "AddCartItemRequest", "CartResDTO", "CartService", "Idempotency-Key tùy chọn", "SỬA"],
    ["CART-03", "Cart", "PUT", "/api/cart/items/{variantId}", "USER/GUEST", "Đổi quantity hoặc xóa khi policy cho phép.", "UpdateCartItemRequest", "CartResDTO", "CartService", "-", "SỬA"],
    ["CART-04", "Cart", "DELETE", "/api/cart/items/{variantId}", "USER/GUEST", "Xóa dòng khỏi đúng giỏ của actor.", "path", "204", "CartService", "-", "GIỮ"],
    ["CO-01", "Checkout", "POST", "/api/checkout", "USER", "Checkout các item được chọn, áp voucher, snapshot và tạo PENDING_APPROVAL; chưa tạo payment.", "CheckoutRequest", "CheckoutResponseDTO", "CheckoutService", "Bắt buộc Idempotency-Key", "TẠO MỚI"],
    ["CO-02", "Checkout", "POST", "/api/guest-checkout", "GUEST", "Checkout CATALOG COD, contact bắt buộc, tạo token link và PENDING_APPROVAL.", "GuestCheckoutRequest", "CheckoutResponseDTO", "CheckoutService", "Bắt buộc Idempotency-Key", "SỬA"],
    ["CR-01", "WF01", "POST", "/api/custom-requests", "USER", "Tạo một draft custom mới.", "CreateCustomRequest", "CustomRequestDTO", "CustomRequestService", "Idempotency-Key khuyến nghị", "TẠO MỚI"],
    ["CR-02", "WF01", "GET", "/api/custom-requests", "USER", "Danh sách các draft/yêu cầu custom của chính USER.", "page/status", "CustomRequestDTO page", "CustomRequestService", "-", "TẠO MỚI"],
    ["CR-03", "WF01", "GET", "/api/custom-requests/{id}", "USER", "Xem draft của chính mình.", "path", "CustomRequestDTO", "CustomRequestService", "-", "TẠO MỚI"],
    ["CR-04", "WF01", "PUT", "/api/custom-requests/{id}", "USER", "Cập nhật draft chưa submit.", "UpdateCustomRequest", "CustomRequestDTO", "CustomRequestService", "expectedVersion", "TẠO MỚI"],
    ["CR-05", "WF01", "DELETE", "/api/custom-requests/{id}", "USER", "Xóa draft chưa submit.", "path + expectedVersion", "204", "CustomRequestService", "Idempotency-Key", "TẠO MỚI"],
    ["CR-06", "WF01", "POST", "/api/custom-requests/{id}/submit", "USER", "Chụp snapshot draft thành CUSTOM Order PENDING_APPROVAL đúng một lần.", "SubmitCustomRequest", "CheckoutResponseDTO", "CustomRequestService", "Bắt buộc Idempotency-Key", "TẠO MỚI"],
    ["ORD-01", "Order query", "GET", "/api/orders/me", "USER", "Danh sách Order của USER theo filter trạng thái.", "page/status/type", "OrderListDTO page", "OrderQueryService", "-", "SỬA"],
    ["ORD-02", "Order query", "GET", "/api/orders/{id}", "USER/STAFF/MANAGER", "Xem Order theo access policy; không suy quyền chat.", "path", "OrderDTO", "OrderQueryService", "-", "SỬA"],
    ["ORD-03", "Order query", "GET", "/api/orders/{id}/history", "USER/STAFF/MANAGER", "Xem lịch sử nghiệp vụ đã được phép hiển thị.", "path/page", "OrderStatusHistoryDTO page", "OrderQueryService", "-", "SỬA"],
    ["MGR-01", "WF02", "GET", "/api/manager/orders", "MANAGER", "Danh sách đơn theo trạng thái/loại để review.", "page/status/type", "OrderListDTO page", "OrderQueryService", "-", "TẠO MỚI"],
    ["MGR-02", "WF02", "POST", "/api/manager/orders/{id}/review", "MANAGER", "Duyệt khả năng làm hoặc từ chối; có thể chọn staff hợp lệ.", "ReviewOrderRequest", "OrderActionResponse", "ManagerOrderService", "Idempotency-Key + expectedVersion", "SỬA"],
    ["MGR-03", "WF03", "GET", "/api/manager/staff/available", "MANAGER", "Danh sách tham khảo staff AVAILABLE chưa có active assignment.", "page", "StaffSummaryDTO page", "OrderAssignmentService", "-", "TẠO MỚI"],
    ["MGR-04", "WF03", "POST", "/api/manager/orders/{id}/assignment", "MANAGER", "Phân staff nguyên tử; manager không trở thành member chat.", "AssignOrderRequest", "OrderAssignmentDTO", "OrderAssignmentService", "Idempotency-Key + expectedVersion", "SỬA"],
    ["STF-01", "WF03", "GET", "/api/staff/orders", "STAFF", "Danh sách đơn của staff và đơn có thể claim theo filter.", "page/status/view", "OrderListDTO page", "OrderQueryService", "-", "SỬA"],
    ["STF-02", "WF03", "POST", "/api/staff/orders/{id}/claim", "STAFF", "Tự nhận Order nếu AVAILABLE và chưa giữ đơn khác.", "ClaimOrderRequest", "OrderAssignmentDTO", "OrderAssignmentService", "Idempotency-Key + expectedVersion", "SỬA"],
    ["CHAT-01", "WF03", "GET", "/api/orders/{id}/chat", "OWNER USER/ASSIGNED STAFF", "Lấy thông tin thread riêng của Order.", "path", "ChatThreadDTO", "OrderChatService", "-", "TẠO MỚI"],
    ["CHAT-02", "WF03", "GET", "/api/orders/{id}/chat/messages", "OWNER USER/ASSIGNED STAFF", "Đọc lịch sử message phân trang.", "cursor/page", "OrderChatMessageDTO page", "OrderChatService", "-", "TẠO MỚI"],
    ["CHAT-03", "WF03", "WS", "/app/orders/{id}/chat/messages", "OWNER USER/ASSIGNED STAFF", "Gửi message realtime sau kiểm tra membership.", "SendOrderMessageRequest", "OrderChatMessageDTO", "OrderChatService", "clientMessageId", "TẠO MỚI"],
    ["AGR-01", "WF03", "POST", "/api/orders/{id}/agreements", "OWNER USER", "Gửi form thỏa thuận; form hợp lệ dừng timer custom ngay.", "SubmitAgreementRequest", "OrderAgreementDTO", "OrderAgreementService", "Idempotency-Key + expectedVersion", "TẠO MỚI"],
    ["AGR-02", "WF03", "GET", "/api/orders/{id}/agreements", "OWNER USER/ASSIGNED STAFF", "Xem lịch sử phiên bản agreement.", "page", "OrderAgreementDTO page", "OrderAgreementService", "-", "TẠO MỚI"],
    ["AGR-03", "WF03", "POST", "/api/orders/{id}/agreements/{agreementId}/review", "ASSIGNED STAFF", "Xác nhận đúng thông tin hoặc yêu cầu sửa; accepted -> ORDER_ACCEPTED.", "ReviewAgreementRequest", "OrderActionResponse", "OrderAgreementService", "Idempotency-Key + versions", "TẠO MỚI"],
    ["PAY-01", "WF04", "POST", "/api/orders/{id}/payment-attempts", "OWNER USER", "Mở thanh toán ONLINE sau ORDER_ACCEPTED, hạn 1 giờ.", "CreatePaymentAttemptRequest", "PaymentAttemptDTO", "OrderPaymentService", "Bắt buộc Idempotency-Key", "TẠO MỚI"],
    ["PAY-02", "WF04", "GET", "/api/orders/{id}/payment-attempts/{attemptId}", "OWNER USER/ASSIGNED STAFF", "Đọc trạng thái attempt đã được lọc dữ liệu nhạy cảm.", "path", "PaymentAttemptDTO", "OrderPaymentService", "-", "TẠO MỚI"],
    ["PAY-03", "WF04", "POST", "/api/payment/momo-callback", "PAYMENT PROVIDER", "Webhook xác thực chữ ký/số tiền, xử lý trùng, ghi PAID và báo staff.", "MomoWebhookRequest", "Provider ACK", "PaymentWebhookService", "provider transaction unique", "SỬA"],
    ["PROD-01", "WF04", "POST", "/api/staff/orders/{id}/production/start", "ASSIGNED STAFF", "Staff chủ động bắt đầu; ORDER_ACCEPTED -> ORDER_CREATING và tính deadline.", "StartProductionRequest", "OrderActionResponse", "ProductionService", "Idempotency-Key + expectedVersion", "TẠO MỚI"],
    ["CP-01", "WF05", "GET", "/api/orders/{id}/checkpoints", "OWNER USER/ASSIGNED STAFF/MANAGER", "Đọc checkpoint theo quyền; USER chỉ thấy dữ liệu được công bố.", "page", "ProductionCheckpointDTO page", "ProductionCheckpointService", "-", "TẠO MỚI"],
    ["CP-02", "WF05", "POST", "/api/staff/orders/{id}/items/{itemId}/checkpoints", "ASSIGNED STAFF", "Nộp checkpoint; INITIAL_SHAPE gửi thẳng báo cáo tiến trình.", "SubmitCheckpointRequest", "ProductionCheckpointDTO", "ProductionCheckpointService", "Idempotency-Key", "TẠO MỚI"],
    ["CP-03", "WF05", "POST", "/api/manager/orders/{id}/checkpoints/{checkpointId}/review", "MANAGER", "Review FINAL_QC; qua mọi loại -> READY_TO_SHIP và release staff.", "ReviewCheckpointRequest", "OrderActionResponse", "ProductionCheckpointService", "Idempotency-Key + version", "SỬA"],
    ["CHG-01", "WF06", "POST", "/api/orders/{id}/change-requests", "OWNER USER", "Xin đổi chi tiết chưa làm, không được đổi tiền/deadline.", "CreateChangeRequest", "ChangeRequestDTO", "ChangeRequestService", "Idempotency-Key + version", "TẠO MỚI"],
    ["CHG-02", "WF06", "GET", "/api/orders/{id}/change-requests", "OWNER USER/ASSIGNED STAFF", "Xem các yêu cầu đổi của Order.", "page", "ChangeRequestDTO page", "ChangeRequestService", "-", "TẠO MỚI"],
    ["CHG-03", "WF06", "POST", "/api/orders/{id}/change-requests/{requestId}/review", "ASSIGNED STAFF", "Chấp nhận hoặc từ chối; từ chối không dừng sản xuất.", "ReviewChangeRequest", "ChangeRequestDTO", "ChangeRequestService", "Idempotency-Key + version", "TẠO MỚI"],
    ["SHIP-01", "WF07", "POST", "/api/staff/orders/{id}/shipment", "HISTORIC ASSIGNED STAFF", "Ghi exportedQuantity cho mọi item và READY_TO_SHIP -> SHIPPING.", "ShipOrderRequest", "OrderActionResponse", "OrderShippingService", "Idempotency-Key + expectedVersion", "SỬA"],
    ["REC-01", "WF08", "POST", "/api/orders/{id}/receipt", "OWNER USER", "Ghi số nhận, chấp nhận chênh lệch và chuyển DELIVERED.", "ReceiptConfirmRequest", "OrderReceiptDTO", "OrderReceiptService", "Idempotency-Key + expectedVersion", "SỬA"],
    ["GST-01", "WF08", "GET", "/api/guest/orders/{id}", "GUEST TOKEN READ", "Tra cứu Order guest bằng token có scope/expiry/rate-limit.", "X-Guest-Order-Token", "OrderDTO", "OrderQueryService", "token scope READ", "TẠO MỚI"],
    ["GST-02", "WF08", "POST", "/api/guest/orders/{id}/receipt", "GUEST TOKEN CONFIRM_RECEIPT", "Xác nhận nhận hàng như USER, không chỉ dựa orderId.", "ReceiptConfirmRequest + token", "OrderReceiptDTO", "OrderReceiptService", "Idempotency-Key + token", "TẠO MỚI"],
    ["CAN-01", "WF09", "POST", "/api/orders/{id}/cancel", "OWNER USER", "Hủy trước ORDER_ACCEPTED, lưu lý do và gọi kernel hủy.", "CancelOrderRequest", "OrderActionResponse", "CancelOrderService", "Idempotency-Key + expectedVersion", "SỬA"],
    ["CAN-02", "WF09", "POST", "/api/guest/orders/{id}/cancel", "GUEST TOKEN CANCEL", "Guest hủy trước ORDER_ACCEPTED; hoàn quota nhưng giữ anti-spam.", "CancelOrderRequest + token", "OrderActionResponse", "CancelOrderService", "Idempotency-Key + token", "TẠO MỚI"],
    ["REO-01", "WF08", "POST", "/api/orders/{id}/reorder", "OWNER USER", "Catalog thêm item hợp lệ vào giỏ; custom tạo draft mới.", "ReorderRequest", "ReorderResponseDTO", "ReorderService", "Idempotency-Key", "SỬA"],
]


service_rows = [r for r in objects if r[0] == "Service"]
entity_rows = [r for r in objects if r[0] == "Entity"]
repository_rows = [r for r in objects if r[0] == "Repository"]
controller_rows = [r for r in objects if r[0] == "Controller"]


entity_fields = {
    "Order": [
        ["id", "Long", "Y", "PK", "Định danh Order"], ["version", "Long", "Y", "@Version", "Khóa lạc quan"],
        ["orderType", "OrderType", "Y", "CATALOG/CUSTOM", "Loại nguồn đơn"], ["status", "OrderStatus", "Y", "PENDING_APPROVAL", "Lifecycle chính"],
        ["paymentMethodType", "PaymentMethod", "Y", "COD/ONLINE", "Phương thức chuẩn hóa"], ["paymentStatus", "PaymentStatus", "N", "NOT_DUE/PENDING/PAID/...", "Legacy chưa rõ giữ null"],
        ["userId / user", "Long/User", "N", "null với guest", "Chủ tài khoản"], ["guestSessionId", "String", "N", "guest only", "Không phải quyền tra cứu Order"],
        ["contactSnapshot", "Embedded", "Y", "server snapshot", "Liên hệ tại checkout"], ["assignedStaffId", "Long", "N", "history retained", "Staff đã phụ trách"],
        ["totalPrice", "double", "Y", ">=0", "Tổng trước voucher"], ["discountAmount", "Double", "Y", ">=0", "Giảm giá"], ["finalPrice", "Double", "N", "CUSTOM pre-agreement=null", "Giá cuối"],
        ["managerApprovedAt", "LocalDateTime", "N", "server time", "Manager duyệt"], ["confirmationDueAt", "LocalDateTime", "N", "+24h CUSTOM", "Hạn gửi form"],
        ["productionStartedAt", "LocalDateTime", "N", "ORDER_CREATING", "Mốc tính deadline"], ["readyToShipAt", "LocalDateTime", "N", "checkpoint 2 pass", "Sẵn sàng giao"],
        ["cancelledAt/reason/source", "metadata", "N", "CANCELLED", "Audit hủy"], ["guestTokenHash/scopes/expiry/revoked", "security", "N", "guest only", "Quyền link email"],
    ],
    "OrderItem": [
        ["id", "Long", "Y", "PK", "Định danh dòng hàng"], ["orderId/order", "Long/Order", "Y", "FK", "Order cha"],
        ["sourceType", "OrderItemSourceType", "Y", "CATALOG/CUSTOM", "Nguồn item"], ["quantity", "int", "Y", ">0", "Số lượng đặt"],
        ["price", "Double", "N", "CUSTOM chưa chốt=null", "Đơn giá snapshot"], ["productNameSnapshot", "String", "N", "catalog", "Tên sản phẩm tại checkout"],
        ["variantNameSnapshot", "String", "N", "catalog", "Tên biến thể"], ["specSnapshot", "TEXT/JSON", "Y", "immutable snapshot", "Thông số catalog/custom"],
        ["madeDaySnapshot", "Double", "N", "catalog >=2", "madeDay tại checkout"], ["productionDurationDays", "Integer", "N", "computed", "Số ngày sản xuất item"],
        ["computedCompletionAt", "LocalDateTime", "N", "at ORDER_CREATING", "Hạn item"], ["durationRuleVersion", "String", "N", "versioned", "Truy vết công thức"],
        ["exportedQuantity", "Integer", "N", "0..quantity", "Bàn giao"], ["receivedQuantity", "Integer", "N", "0..exported", "Thực nhận"],
    ],
    "CustomRequest": [["id", "Long", "Y", "PK", "Draft"], ["version", "Long", "Y", "@Version", "Concurrency"], ["userId", "Long", "Y", "owner", "Chỉ USER"], ["status", "CustomRequestStatus", "Y", "DRAFT/SUBMITTED/...", "Vòng đời request"], ["name/description/spec", "String/TEXT", "Y", "validated", "Thông tin custom"], ["quantity", "int", "Y", ">0", "Số lượng"], ["attachments", "JSON/ref", "N", "safe URLs", "Ảnh yêu cầu"], ["submittedOrderId", "Long", "N", "unique after submit", "Chống tạo nhiều Order"]],
    "OrderAssignment": [["id", "Long", "Y", "PK", "Assignment"], ["orderId", "Long", "Y", "unique active", "Order"], ["staffId", "Long", "Y", "unique active", "Staff"], ["source", "AssignmentSource", "Y", "MANAGER_ASSIGN/SELF_CLAIM", "Nguồn"], ["status", "AssignmentStatus", "Y", "ACTIVE/RELEASED", "Tình trạng"], ["activeOrderId/activeStaffId", "Long", "N", "DB unique", "Khóa nguyên tử"], ["assignedAt/releasedAt", "LocalDateTime", "Y/N", "server time", "Audit"], ["releaseReason", "AssignmentReleaseReason", "N", "QC_PASS/CANCELLED", "Lý do giải phóng"]],
    "OrderAgreement": [["id", "Long", "Y", "PK", "Agreement"], ["orderId", "Long", "Y", "FK", "Order"], ["agreementVersion", "Integer", "Y", "unique/order", "Phiên bản nghiệp vụ"], ["status", "OrderAgreementStatus", "Y", "DRAFT/SUBMITTED/CONFIRMED/REVISION", "Trạng thái"], ["itemSnapshot", "TEXT", "Y", "immutable", "Spec/quantity"], ["subtotal/discount/finalPrice", "Double", "N", "custom pre-confirm=null", "Giá đã trao đổi"], ["deadlineDays", "Integer", "N", ">0", "Số ngày chốt"], ["submittedBy/At", "String/time", "N", "owner user", "Người gửi"], ["confirmedBy/At", "String/time", "N", "assigned staff", "Người xác nhận"], ["decisionReason", "TEXT", "N", "required on revision", "Lý do"]],
    "PaymentAttempt": [["id", "Long", "Y", "PK", "Attempt"], ["version", "Long", "Y", "@Version", "Concurrency"], ["orderId", "Long", "Y", "FK", "Order"], ["provider", "String", "Y", "MOMO", "Nhà cung cấp"], ["requestReference", "String", "Y", "unique/provider", "Idempotency request"], ["providerTransactionId", "String", "N", "unique/provider", "Id giao dịch"], ["expectedAmount", "Double", "Y", "=Order.finalPrice", "Tiền kỳ vọng"], ["status", "PaymentAttemptStatus", "Y", "PENDING/PAID/FAILED/EXPIRED", "Kết quả"], ["openedAt/dueAt/paidAt", "LocalDateTime", "Y/Y/N", "due=+1h", "Mốc payment"], ["providerResult", "TEXT", "N", "protected", "Payload phục vụ audit"], ["reconciliationRequired", "boolean", "Y", "false", "Callback muộn/xung đột"]],
    "ChatThread & ChatMessage": [["ChatThread.id", "Long/String", "Y", "PK", "Thread"], ["ChatThread.orderId", "Long", "Y", "unique", "Một phòng/Order"], ["ChatThread.userId", "Long", "Y", "owner", "USER thành viên"], ["ChatThread.staffId", "Long", "Y", "assigned", "STAFF thành viên"], ["ChatThread.status", "enum", "Y", "ACTIVE/CLOSED", "Không có manager member"], ["ChatMessage.id", "String", "Y", "PK", "Message"], ["ChatMessage.threadId", "ref", "Y", "FK/ref", "Không query chỉ bằng userId"], ["ChatMessage.senderId", "Long", "Y", "member", "Người gửi"], ["ChatMessage.clientMessageId", "String", "Y", "unique/sender", "Chống gửi trùng"], ["content/attachments", "TEXT/JSON", "N", "validated", "Nội dung"], ["sentAt", "Instant", "Y", "server time", "Thứ tự"]],
    "Production": [["checkpoint.id/version", "Long", "Y", "PK/@Version", "Checkpoint"], ["checkpoint.orderItemId", "Long", "Y", "FK", "Theo loại item"], ["checkpoint.type", "enum", "Y", "INITIAL_SHAPE/FINAL_QC", "Mốc"], ["checkpoint.attempt", "int", "Y", ">0", "Lần nộp/rework"], ["checkpoint.status", "enum", "Y", "SUBMITTED/PASSED/REWORK", "Kết quả"], ["checkpoint.note/submittedAt", "TEXT/time", "N/Y", "server time", "Báo cáo"], ["image.checkpointId/url", "ref/String", "Y", "safe URL", "Ảnh"], ["decision.checkpointId", "Long", "Y", "FK", "Chỉ FINAL_QC"], ["decision.result/reason", "enum/TEXT", "Y/N", "reason on rework", "KCS"]],
    "ChangeRequest & OrderReceipt": [["change.id/version", "Long", "Y", "PK/@Version", "Request"], ["change.orderId/orderItemId", "Long", "Y", "FK", "Phạm vi"], ["change.requestedChanges", "TEXT", "Y", "unimplemented details", "Không price/deadline"], ["change.status", "enum", "Y", "PENDING/ACCEPTED/REJECTED", "Quyết định"], ["change.decisionReason", "TEXT", "N", "on reject", "Lý do"], ["receipt.id/version", "Long", "Y", "PK/@Version", "Receipt"], ["receipt.orderId", "Long", "Y", "unique", "Một xác nhận cuối"], ["receipt.itemsSnapshot", "TEXT/JSON", "Y", "item quantities", "Số lượng nhận"], ["receipt.mismatchAccepted", "boolean", "Y", "explicit", "Chấp nhận chênh lệch"], ["receipt.confirmedBy/At", "String/time", "Y", "USER/guest", "Audit"]],
}


camunda_rows = [
    ["order-lifecycle.bpmn", "TẠO MỚI", "WF02-WF05", "Một process/Order, businessKey=orderId", "Review -> assignment/agreement -> payment -> staff start -> checkpoint/KCS", "DB là nguồn trạng thái chính"],
    ["OrderWorkflowGateway.java", "TẠO MỚI", "WF02-WF05", "Adapter start/complete/correlate", "Không chứa business rule", "Gọi sau commit hoặc qua outbox"],
    ["EnsureOrderChatDelegate.java", "TẠO MỚI", "WF03", "Service task tạo chat", "Idempotent theo orderId", "Gọi OrderChatService"],
    ["CustomConfirmationTimeoutDelegate.java", "TẠO MỚI", "WF03/WF09", "Timer PT24H", "Kiểm tra DB rồi gọi CancelOrderService", "Form hợp lệ đã lưu thì không hủy"],
    ["PaymentTimeoutDelegate.java", "TẠO MỚI", "WF04/WF09", "Timer PT1H", "Khóa Order/Attempt, không hồi sinh callback muộn", "Gọi CancelOrderService/reconciliation"],
    ["OrderWorkflowEventHandler.java", "TẠO MỚI", "WF02-WF05", "Nhận event đã commit", "Correlate message chống trùng", "DurableRequestExecutor"],
    ["approve-cart.bpmn", "MIGRATE", "Legacy", "Không dùng cho đơn mới sau cutover", "Có payment trước review và nghiệp vụ stock cũ", "Kế hoạch instance migration"],
    ["guest-purchase-runtime.bpmn", "MIGRATE", "Legacy", "Không tạo lifecycle cạnh tranh", "Guest dùng cùng lifecycle Order", "Kế hoạch instance migration"],
]

event_rows = [
    ["ORDER_CREATED", "Order", "Checkout/Custom submit", "USER: email+notification; Guest: email", "Mốc PENDING_APPROVAL", "orderId+version"],
    ["ORDER_ACCEPTED", "Order", "Agreement/guest catalog assignment", "USER: email+notification; Guest: email", "Mốc ORDER_ACCEPTED", "orderId+version"],
    ["ORDER_ASSIGNMENT_CREATED", "Assignment", "Assign/claim", "Notification STAFF; tạo chat khi phù hợp", "Không cho manager vào chat", "assignmentId"],
    ["ORDER_PAYMENT_CONFIRMED", "PaymentAttempt", "Webhook PAID", "Notification STAFF phụ trách", "Không tự ORDER_CREATING", "attemptId/providerTxn"],
    ["ORDER_PRODUCTION_STARTED", "Order", "Staff start", "Email owner; notification USER", "Mốc ORDER_CREATING + deadline", "orderId+version"],
    ["ORDER_CHECKPOINT_REPORTED", "Checkpoint", "INITIAL_SHAPE", "Email tiến trình; notification USER 'Đã có hình ảnh đầu tiên của item #...'", "Không đổi Order status", "checkpointId/attempt"],
    ["ORDER_READY_TO_SHIP", "Order", "All FINAL_QC passed", "Email owner; notification USER", "Mốc READY_TO_SHIP", "orderId+version"],
    ["ORDER_CANCELLED", "Order", "Cancel kernel", "Email owner; notification USER", "Thông báo riêng, có reason", "orderId+version"],
]

wf_rows = [
    ["WF01", "Custom request", "USER", "CustomRequestController/Service", "CustomRequest, Order, OrderItem", "PENDING_APPROVAL", "05"],
    ["WF02", "Manager review", "MANAGER", "ManagerOrderController/Service", "Order", "PENDING_ASSIGNMENT/DISCUSSING/ORDER_ACCEPTED/CANCELLED", "07"],
    ["WF03", "Assignment, chat, agreement", "USER/STAFF/MANAGER(assign only)", "Assignment/Chat/Agreement services", "Assignment, ChatThread, Agreement", "ORDER_ACCEPTED", "07-08"],
    ["WF04", "Payment và bắt đầu", "USER/STAFF/Provider", "Payment/Production services", "PaymentAttempt, Order", "ORDER_CREATING", "09"],
    ["WF05", "Checkpoint/KCS", "STAFF/MANAGER", "ProductionCheckpointService", "Checkpoint/Image/Decision", "READY_TO_SHIP", "10"],
    ["WF06", "ChangeRequest", "USER/STAFF", "ChangeRequestService", "ChangeRequest", "Order vẫn ORDER_CREATING", "11"],
    ["WF07", "Bàn giao", "Historic assigned STAFF", "OrderShippingService", "OrderItem", "SHIPPING", "12"],
    ["WF08", "Nhận hàng/reorder", "USER/GUEST", "Receipt/Reorder services", "OrderReceipt/Cart/CustomRequest", "DELIVERED hoặc draft/cart", "12"],
    ["WF09", "Hủy", "USER/GUEST/MANAGER/SYSTEM reason", "CancelOrderService", "Order, Voucher, Assignment", "CANCELLED", "06"],
]

plan_rows = [
    ["01", "Baseline/gap", "VERIFIED", "Không sửa trong tài liệu này", "Báo cáo baseline"],
    ["02", "Domain + Liquibase expand", "VERIFIED", "Giữ changeset đã chạy bất biến", "Entity/enums/migration"],
    ["03", "Consistency infrastructure", "VERIFIED", "Dùng cho mọi command mới", "Transition/idempotency/outbox/token"],
    ["04", "Catalog/cart/checkout", "NEXT", "Triển khai sau khi chốt workbook", "Duration, snapshot, checkout PENDING_APPROVAL"],
    ["05", "Custom request", "NOT_STARTED", "Sau 04", "Draft/submit"],
    ["06", "Cancel kernel", "NOT_STARTED", "Trước review/timer", "CancelOrderService"],
    ["07", "Review/assignment", "NOT_STARTED", "Sau 04-06", "WF02"],
    ["08", "Chat/agreement/timer", "NOT_STARTED", "Sau 05-07", "WF03"],
    ["09", "Payment/start", "NOT_STARTED", "Sau 06-08", "WF04"],
    ["10", "Checkpoint/KCS", "NOT_STARTED", "Sau 07-09", "WF05"],
    ["11", "ChangeRequest", "NOT_STARTED", "Sau 08-10", "WF06"],
    ["12", "Shipping/receipt/reorder", "NOT_STARTED", "Sau 06,10,11", "WF07-WF08"],
    ["13", "Integration/cutover", "NOT_STARTED", "Cuối", "BPMN/events/E2E/legacy removal"],
]

decision_rows = [
    ["D01", "Tên file Java", "Một public top-level type/file; filename trùng chính xác tên class/interface/record.", "ĐÃ CHỐT", "Áp dụng toàn bộ code mới"],
    ["D02", "Tiền", "Dùng double/Double theo quyết định dự án; validate finite/non-negative; custom chưa chốt là null.", "ĐÃ CHỐT", "Không tự đổi BigDecimal"],
    ["D03", "Migration", "Liquibase; changeset cũ không sửa sau khi áp dụng.", "ĐÃ CHỐT", "Thêm file changelog mới"],
    ["D04", "Chat", "Chỉ USER sở hữu và STAFF phụ trách; manager không đọc/gửi/subscribe.", "ĐÃ CHỐT", "AccessPolicy riêng"],
    ["D05", "Camunda", "Điều phối WF02-WF05; DB source of truth; một process/Order.", "ĐÃ CHỐT", "businessKey=orderId"],
    ["D06", "Email trạng thái", "Bốn mốc: PENDING_APPROVAL, ORDER_ACCEPTED, ORDER_CREATING, READY_TO_SHIP.", "ĐÃ CHỐT", "Checkpoint/cancel là mail riêng"],
    ["D07", "Guest token TTL", "Thời hạn cụ thể, cấp lại link và guest refresh session.", "CẦN CHỐT TRƯỚC API GUEST", "Guard đã yêu cầu expiry"],
    ["D08", "Bàn giao thiếu", "Có bắt buộc exportedQuantity=ordered quantity hay cho 0..quantity.", "CẦN CHỐT TRƯỚC WF07", "WF07 hiện mô tả khoảng 0..quantity"],
]

test_rows = [
    ["T04-01", "Duration", "madeDay=3; q=1/5/6/10", "5/14/16/22"],
    ["T04-02", "Calendar", "Bắt đầu thứ Bảy, duration=2, không holiday", "Hạn thứ Ba"],
    ["T04-03", "Checkout", "Chọn một phần giỏ", "Chỉ item chọn bị xóa; snapshot đầy đủ"],
    ["T04-04", "Checkout retry", "Cùng key/cùng payload", "Cùng Order; không dùng voucher lần hai"],
    ["T04-05", "Checkout conflict", "Cùng key/khác payload", "409 REQUEST_KEY_CONFLICT"],
    ["T07-01", "Assignment race", "Hai staff/manager đồng thời", "Một active assignment thắng"],
    ["T08-01", "Chat access", "Manager hoặc staff không phụ trách", "403 và không subscribe WS"],
    ["T08-02", "Timer race", "Submit form và timeout đồng thời", "Chỉ một kết quả commit"],
    ["T09-01", "Webhook duplicate", "Provider gửi lại callback", "PAID/event đúng một lần theo milestone"],
    ["T09-02", "Webhook late", "Callback sau timeout/cancel", "Không hồi sinh Order; reconciliationRequired"],
    ["T10-01", "Checkpoint 1", "Staff gửi ảnh", "Mail tiến trình + notification; không manager review/status change"],
    ["T10-02", "Checkpoint 2", "Mọi loại pass", "READY_TO_SHIP + release assignment cùng transaction"],
    ["T12-01", "Guest access", "Chỉ orderId, thiếu/sai token", "401/403/429 phù hợp; không lộ dữ liệu"],
]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "00_BIA"
    title(ws, "BÁO CÁO THIẾT KẾ CODE VÀ HỢP ĐỒNG ENDPOINT", "Shop handmade made-to-order — bản thống nhất trước khi tiếp tục triển khai", 8)
    section(ws, 4, "Thông tin tài liệu", 8)
    info = [
        ["Tên tài liệu", "Thiết kế file Java, domain object và endpoint tổng hợp"],
        ["Ngày lập", datetime.now().strftime("%d/%m/%Y")],
        ["Nguồn ưu tiên", "Quyết định mới nhất của chủ dự án → docs/target.txt → WF01–WF09 → kế hoạch triển khai → code legacy"],
        ["File mẫu tham khảo", r"C:\Users\Admin\Downloads\ePAS_999_DES_Field & Table Definitions_DGM_v1.0 (3).xlsb"],
        ["Quy ước tên", "Mỗi public top-level Java type nằm trong một file cùng tên đối tượng, ví dụ Order.java chứa Order."],
        ["Phạm vi", "Thiết kế đích. Nhãn ĐÃ CÓ KHUNG không có nghĩa endpoint/nghiệp vụ đã hoàn thành."],
    ]
    table(ws, 5, ["Thuộc tính", "Nội dung"], info, "T_BiaInfo", [25, 110])
    section(ws, 13, "Kết luận kiến trúc", 8)
    paragraph(ws, 14, "Backend giữ mô hình modular monolith. Controller chỉ tiếp nhận HTTP/WS và lấy actor từ JWT/token; service quyết định nghiệp vụ và transaction; repository chỉ truy cập/khóa dữ liệu; Camunda điều phối WF02–WF05 nhưng database là nguồn trạng thái chính; outbox chịu trách nhiệm phát event bền vững.", 8, LIGHT_BLUE)
    paragraph(ws, 15, "Các API command được đặt theo hành động (review, claim, start, cancel...) thay vì cho client tự cập nhật status. Mọi action nhạy cảm kiểm tra actor, trạng thái, expectedVersion và idempotency trong backend.", 8, GREEN)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 110

    toc = wb.create_sheet("01_MUC_LUC")
    title(toc, "MỤC LỤC VÀ CÁCH ĐỌC", "Các sheet danh mục tương ứng cách tổ chức Tables/detail trong file mẫu ePAS", 5)
    toc_rows = [
        ["02_QUY_UOC", "Quy tắc đặt tên, trách nhiệm lớp và nhãn thay đổi", "Đọc trước"],
        ["03_OBJECT_CATALOG", "Toàn bộ file Java liên quan cần giữ/sửa/tạo/migrate", "Master catalog"],
        ["04_ENTITY_CATALOG", "Danh mục entity và vai trò lưu trữ", "Domain"],
        ["E_*", "Định nghĩa trường mục tiêu của entity trọng tâm", "Entity detail"],
        ["05_DTO_CATALOG", "Request/response contract cần tạo hoặc sửa", "API contract"],
        ["06_SERVICE_CATALOG", "Service và ranh giới trách nhiệm", "Application layer"],
        ["07_REPOSITORY", "Repository/lock/query cần có", "Persistence"],
        ["08_CONTROLLER", "Controller và phạm vi route", "Web layer"],
        ["09_ENDPOINTS", "Endpoint tổng hợp: dùng làm gì, ai gọi, DTO/service", "API master"],
        ["10_CAMUNDA_EVENT", "BPMN, delegate và event/notification", "Workflow"],
        ["11_WF_MAPPING", "Ánh xạ WF01–WF09 sang code/data/state", "Traceability"],
        ["12_KE_HOACH", "Thứ tự triển khai 13 phần", "Delivery"],
        ["13_QUYET_DINH", "Điểm đã chốt và còn mở", "Governance"],
        ["14_KIEM_THU", "Kịch bản kiểm thử chấp nhận", "Quality"],
    ]
    table(toc, 4, ["Sheet", "Nội dung", "Vai trò"], toc_rows, "T_Toc", [30, 90, 25])
    toc.sheet_view.showGridLines = False

    convention_rows = [
        ["File/type", "Một public top-level class/interface/record mỗi file; filename trùng tên type.", "OrderAssignmentService.java chứa OrderAssignmentService"],
        ["Controller", "Không chứa transaction/business rule; lấy actor từ security context; validate request; gọi đúng một application service.", "Không nhận userId để tự nhận quyền"],
        ["Service", "Một capability nghiệp vụ rõ ràng; kiểm tra quyền, trạng thái, version; sở hữu transaction.", "CheckoutService, CancelOrderService"],
        ["Repository", "Truy vấn, khóa và persistence; không gửi email, complete Camunda hay đổi state ngoài service.", "findByIdForUpdate"],
        ["DTO request", "Không cho client gửi status/paid/assigned actor tùy ý; command có expectedVersion khi cần.", "ReviewOrderRequest"],
        ["DTO response", "Không trả entity trực tiếp; ẩn token hash/provider payload; custom price chưa chốt là null.", "PaymentAttemptDTO"],
        ["Mapper", "Chỉ chuyển dữ liệu; không tính tiền, không đổi trạng thái, không gọi repository.", "OrderMapper"],
        ["Camunda delegate", "Adapter mỏng gọi service idempotent; không sao chép business rule.", "PaymentTimeoutDelegate"],
        ["Event", "Ghi outbox cùng transaction; consumer idempotent theo eventId/milestone.", "ORDER_CREATED"],
        ["GIỮ", "Cấu trúc/hành vi chính phù hợp; có thể bổ sung test nhỏ.", "Màu xanh"],
        ["SỬA", "File tồn tại nhưng contract/nghiệp vụ phải đổi.", "Màu vàng"],
        ["TẠO MỚI", "Chưa có file/type mục tiêu.", "Màu xanh dương nhạt"],
        ["ĐÃ CÓ KHUNG", "Entity/repository đã có skeleton từ phần 02; endpoint/service nghiệp vụ chưa hoàn thành.", "Màu cam"],
        ["MIGRATE", "Tồn tại nhưng cần chuyển dần sang mô hình mới.", "Màu cam"],
        ["LOẠI SAU MIGRATION", "Ngoài target; chỉ xóa khi không còn dữ liệu/consumer phụ thuộc.", "Màu đỏ"],
    ]
    add_sheet(wb, "02_QUY_UOC", "QUY ƯỚC THIẾT KẾ VÀ ĐẶT TÊN", "Áp dụng bắt buộc cho các file mới và khi tách file legacy", ["Đối tượng", "Quy tắc", "Ví dụ/Ghi chú"], convention_rows, [24, 95, 45])

    add_sheet(wb, "03_OBJECT_CATALOG", "DANH MỤC FILE JAVA TỔNG HỢP", "Master catalog: lọc theo loại, trạng thái, phase hoặc workflow", ["Loại", "Tên file", "Package", "Phân loại", "Phần", "Workflow", "Mục đích", "Phụ thuộc chính"], objects, [15, 38, 18, 22, 14, 18, 80, 35])
    add_sheet(wb, "04_ENTITY_CATALOG", "DANH MỤC ENTITY", "Entity target và chiến lược giữ/sửa/migrate", ["Loại", "Tên file", "Package", "Phân loại", "Phần", "Workflow", "Mục đích", "Phụ thuộc chính"], entity_rows, [14, 35, 16, 22, 14, 18, 80, 30])

    for sheet_name, entity_name in [
        ("E_Order", "Order"), ("E_OrderItem", "OrderItem"), ("E_CustomRequest", "CustomRequest"),
        ("E_Assignment", "OrderAssignment"), ("E_Agreement", "OrderAgreement"), ("E_Payment", "PaymentAttempt"),
        ("E_Chat", "ChatThread & ChatMessage"), ("E_Production", "Production"), ("E_ChangeReceipt", "ChangeRequest & OrderReceipt"),
    ]:
        rows = [[i + 1] + r for i, r in enumerate(entity_fields[entity_name])]
        add_sheet(wb, sheet_name, f"ENTITY DETAIL — {entity_name}", "Cấu trúc trường mục tiêu phục vụ triển khai; trường legacy chỉ giữ khi migration còn cần", ["STT", "Thuộc tính", "Kiểu Java/DB", "Bắt buộc", "Ràng buộc/Mặc định", "Ý nghĩa"], rows, [8, 35, 24, 14, 38, 68])

    add_sheet(wb, "05_DTO_CATALOG", "DANH MỤC DTO", "Request/response tách riêng; không dùng Entity làm API contract", ["Tên file", "Loại", "Phân loại", "Phần", "Trường chính", "Mục đích/Quy tắc"], dto_rows, [40, 15, 20, 12, 70, 80])
    add_sheet(wb, "06_SERVICE_CATALOG", "DANH MỤC SERVICE", "Ranh giới trách nhiệm để OrderService không tiếp tục phình to", ["Loại", "Tên file", "Package", "Phân loại", "Phần", "Workflow", "Mục đích", "Phụ thuộc chính"], service_rows, [14, 42, 16, 20, 12, 18, 85, 36])
    add_sheet(wb, "07_REPOSITORY", "DANH MỤC REPOSITORY", "Các repository mới/sửa và yêu cầu khóa dữ liệu", ["Loại", "Tên file", "Package", "Phân loại", "Phần", "Workflow", "Mục đích", "Phụ thuộc"], repository_rows, [14, 42, 16, 20, 12, 16, 90, 25])
    add_sheet(wb, "08_CONTROLLER", "DANH MỤC CONTROLLER", "Controller chỉ điều phối HTTP; endpoint chi tiết nằm ở sheet 09_ENDPOINTS", ["Loại", "Tên file", "Package", "Phân loại", "Phần", "Workflow", "Mục đích", "Phụ thuộc"], controller_rows, [14, 40, 16, 20, 12, 18, 90, 20])
    add_sheet(wb, "09_ENDPOINTS", "DANH MỤC ENDPOINT TỔNG HỢP", "Mỗi endpoint nêu rõ dùng để làm gì, actor, DTO, service và cơ chế chống trùng", ["Mã", "Nhóm", "Method", "URL đích", "Actor/Quyền", "Định nghĩa tổng quát", "Request", "Response", "Service xử lý", "Concurrency/Idempotency", "Phân loại"], endpoint_rows, [12, 16, 10, 48, 28, 95, 40, 32, 38, 45, 18])

    ws_flow = wb.create_sheet("10_CAMUNDA_EVENT")
    title(ws_flow, "CAMUNDA, DELEGATE VÀ EVENT", "Camunda điều phối WF02–WF05; outbox phát event sau commit", 6)
    section(ws_flow, 4, "A. File BPMN và Java workflow", 6)
    end = table(ws_flow, 5, ["File", "Phân loại", "Workflow", "Vai trò", "Quy tắc", "Phụ thuộc"], camunda_rows, "T_Camunda", [42, 20, 20, 55, 75, 40])
    section(ws_flow, end + 2, "B. Event nghiệp vụ", 6)
    table(ws_flow, end + 3, ["Event", "Aggregate", "Nguồn phát", "Consumer/kênh", "Ý nghĩa", "Dedup key"], event_rows, "T_Events", [38, 20, 38, 70, 65, 32])
    ws_flow.sheet_view.showGridLines = False
    ws_flow.freeze_panes = "A6"

    add_sheet(wb, "11_WF_MAPPING", "ÁNH XẠ WORKFLOW SANG CODE", "Traceability WF01–WF09", ["Workflow", "Nội dung", "Tác nhân", "Service/Controller chính", "Entity chính", "Kết quả", "Phần triển khai"], wf_rows, [12, 38, 34, 65, 55, 55, 18])
    add_sheet(wb, "12_KE_HOACH", "KẾ HOẠCH TRIỂN KHAI", "Thứ tự dependency đã thống nhất; Phần 04 là phần tiếp theo sau khi duyệt báo cáo", ["Phần", "Phạm vi", "Trạng thái", "Điều kiện/Ghi chú", "Đầu ra"], plan_rows, [10, 45, 18, 70, 70])
    add_sheet(wb, "13_QUYET_DINH", "QUYẾT ĐỊNH VÀ ĐIỂM CẦN CHỐT", "Không tự biến đề xuất kỹ thuật chưa duyệt thành policy nghiệp vụ", ["Mã", "Chủ đề", "Quyết định/Nội dung", "Trạng thái", "Tác động"], decision_rows, [10, 28, 100, 28, 60])
    add_sheet(wb, "14_KIEM_THU", "KỊCH BẢN KIỂM THỬ CHẤP NHẬN", "Danh sách tối thiểu; mỗi phase bổ sung unit/integration/concurrency/security test", ["Mã", "Nhóm", "Dữ liệu/Tình huống", "Kết quả bắt buộc"], test_rows, [14, 24, 75, 95])

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
        sheet.oddFooter.center.text = "Thiết kế code và endpoint — made-to-order"
        sheet.oddFooter.right.text = "Page &P / &N"
        sheet.sheet_view.zoomScale = 85

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def verify_workbook():
    wb = load_workbook(OUTPUT, read_only=False, data_only=False)
    required = {"00_BIA", "03_OBJECT_CATALOG", "04_ENTITY_CATALOG", "05_DTO_CATALOG", "06_SERVICE_CATALOG", "08_CONTROLLER", "09_ENDPOINTS", "10_CAMUNDA_EVENT"}
    missing = required.difference(wb.sheetnames)
    if missing:
        raise RuntimeError(f"Thiếu sheet: {sorted(missing)}")
    endpoint_count = wb["09_ENDPOINTS"].max_row - 4
    object_count = wb["03_OBJECT_CATALOG"].max_row - 4
    if endpoint_count < 40 or object_count < 70:
        raise RuntimeError(f"Báo cáo chưa đủ dữ liệu: objects={object_count}, endpoints={endpoint_count}")
    wb.close()
    return object_count, endpoint_count


if __name__ == "__main__":
    build_workbook()
    objects_count, endpoints_count = verify_workbook()
    print(f"Created: {OUTPUT}")
    print(f"Objects: {objects_count}; Endpoints: {endpoints_count}")

